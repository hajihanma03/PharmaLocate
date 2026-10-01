# -*- coding: utf-8 -*-
"""Fill TCP Form 4 (tech adviser) and Form 5 (subject teacher) — same layout, complete cells."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Pt, RGBColor, Emu, Twips
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn as qn2

OUT_DIR = Path(r"c:\Users\ADMIN\Downloads")
GREEN = RGBColor(0x0B, 0x5C, 0x3A)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x44, 0x44, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "0B5C3A"
ALT = "F3F7F5"


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run(run, size=11, bold=False, color=DARK, center=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def cell_para(cell, text, *, bold=False, size=9, color=DARK, center=False, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=WHITE if white else color)


def set_cell_margins(cell, **kw):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in kw.items():
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def header_block(doc, form_no, rev_label, role_name, role_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("TARLAC STATE UNIVERSITY")
    set_run(r, size=12, bold=True, color=GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("COLLEGE OF COMPUTER STUDIES")
    set_run(r, size=11, bold=True, color=GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{form_no}  (Rev. 2)")
    set_run(r, size=14, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("PROGRESS REPORT")
    set_run(r, size=13, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Academic Year 2025–2026  ·  Second Semester")
    set_run(r, size=10, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(role_name)
    set_run(r, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(f"({role_title})")
    set_run(r, size=10, color=MUTED)


def meta_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.autofit = True
    for i, (k, v) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        shade(c0, "E8F0EC")
        cell_para(c0, k, bold=True, size=9)
        cell_para(c1, v, size=9)
        c0.width = Cm(3.6)
        c1.width = Cm(14.2)
    doc.add_paragraph()


def data_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        shade(c, HEADER_BG)
        cell_para(c, h, bold=True, size=8, white=True, center=True)
        c.width = Cm(widths[i])
        set_cell_margins(c, top=40, bottom=40, left=50, right=50)
    for ri, row in enumerate(rows):
        prevent_row_split(table.rows[ri + 1])
        for ci, val in enumerate(row):
            c = table.rows[ri + 1].cells[ci]
            c.width = Cm(widths[ci])
            if ri % 2:
                shade(c, ALT)
            center = ci in (0, len(headers) - 1)
            size = 8 if ci else 8
            cell_para(c, val, size=8, center=center, bold=(ci == 0))
            set_cell_margins(c, top=40, bottom=40, left=50, right=50)
    doc.add_paragraph()
    return table


def sign_block(doc, left_name, left_role, right_name, right_role):
    p = doc.add_paragraph()
    r = p.add_run("Prepared by the proponents:\t\t\tNoted:")
    set_run(r, size=10)

    t = doc.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text = ""
    cell_para(t.rows[0].cells[0], "\n\n______________________________", size=10, center=True)
    cell_para(t.rows[0].cells[1], "\n\n______________________________", size=10, center=True)
    cell_para(t.rows[1].cells[0], left_name, size=9, bold=True, center=True)
    cell_para(t.rows[1].cells[1], right_name, size=9, bold=True, center=True)
    cell_para(t.rows[2].cells[0], left_role, size=8, color=MUTED, center=True)
    cell_para(t.rows[2].cells[1], right_role, size=8, color=MUTED, center=True)


def footer_line(doc, form_sf):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        f"Form No.: {form_sf}     Revision No.: 02     Effectivity Date: September 01, 2017     Page  "
    )
    set_run(r, size=8, color=MUTED)


def make_form4():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.5)
        s.right_margin = Cm(1.5)
        s.page_width = Cm(21.59)
        s.page_height = Cm(27.94)

    header_block(doc, "TCP FORM 4", "Rev. 2", "KWINNO L. PINEDA", "TECHNICAL ADVISER")
    meta_table(
        doc,
        [
            ("Project title", "PharmaLocate: A Web-Based Pharmacy Inquiry System with Geofencing for Real-Time Drug Availability, Sale & Inventory Checking"),
            ("Proponents", "Chua, Calvin O.; Aurelio, Lester C.; Mercado, Arji Chesner G.; Lacanlale, Nathalie Angel C."),
            ("Period covered", "July 2026 – 1 October 2026"),
            ("Workspace", "Programming4 (Laravel 13 + 3-file frontend + MySQL)"),
            ("Purpose of this form", "Record of technical-adviser consultations: section of the study, recommendation given, and status in the current system or manuscript."),
        ],
    )
    p = doc.add_paragraph()
    r = p.add_run(
        "Instructions: DATE is the consultation when the recommendation was given. "
        "SECTION / PART is the chapter or component. RECOMMENDATIONS are the adviser’s instructions. "
        "STATUS is how the team complied as of 1 October 2026. SIGNATURE is left blank for the technical adviser."
    )
    set_run(r, size=9, color=MUTED)

    rows = [
        [
            "Jul 2026",
            "Title",
            "Include POS in the title. Recommended wording: A Web-Based Pharmacy Inquiry System with Geofencing for Real-Time Drug Availability, Sale & Inventory Checking.",
            "Complied in the live UI (browser title, navbar, login subtitle) and for the manuscript title page. POS is a named module (sales + inventory).",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 1 Project Context",
            "Make the problem more powerful and connect it to the SDGs.",
            "For the manuscript: strengthen access-to-medicines / SDG 3 wording. System already operationalizes access via live availability and locator around TPH.",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 1 Introduction",
            "The introduction is too long; shorten it while keeping vital information.",
            "For the manuscript: compress Chapter 1. UI uses a short product subtitle instead of the long former tagline.",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 1 Purpose / Scope",
            "Add the scope. Do not include pharmacies inside Tarlac Provincial Hospital (TPH).",
            "Complied in the system: Home scope banner; public GET /api/pharmacies and /api/availability hide inside_tph stores; admin checkbox “Inside TPH.” Seeded SpaRx and Magic 8 are community pharmacies around TPH. Run php artisan migrate for the column.",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 1 Objectives (1.3)",
            "Revise Objective 2 to proper format. Include inventory management in 1.3.",
            "Complied in the system: Admin sidebar “Inventory management” and “POS — sales.” Align the same labels in manuscript 1.3.",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 1 Scope & Limitation",
            "Present White Box Testing in table form, with parameters broken down.",
            "Test IDs WB-01–WB-12 are tabulated in Documentation_v1.3.md §40 and executed in PHPUnit (Sprint7WhiteBoxTest). Copy that table into Chapter 1 as required.",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 2 Local literature",
            "Remove the researchers’ own study from the local literature (LLM).",
            "Manuscript only: delete self-citation from Chapter 2 LLM.",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 2 Conceptual framework",
            "Revise the conceptual framework to proper IPOO formatting.",
            "Manuscript figure: Input–Process–Output–Outcome. Live mapping: user/device → Laravel API + geofence/POS logic → MySQL → UI (availability, tickets, sales).",
            "",
        ],
        [
            "Jul 2026",
            "Ch. 2 Definition of terms",
            "Add definitions of the system features.",
            "Define terms to match live labels: Inquiries, Pharmacies/geofence locator, Medicines/availability, Inventory management, POS—sales, Geofences, Users, Help assistant.",
            "",
        ],
        [
            "Aug 2026",
            "Ch. 3 Architecture",
            "Block-diagram system architecture; make it per icon.",
            "Manuscript figure with one icon per block. Running stack to label: Browser → HTML/CSS/JS → Laravel 13 → MySQL; optional Tile38; Leaflet/OSM map.",
            "",
        ],
        [
            "Aug 2026",
            "Ch. 3 Sources of data",
            "Connect sources of data to the objectives.",
            "Manuscript table: SpaRx and Magic 8 (seeded/participating stores) mapped to 1.3 (availability, inquiry, geofence, inventory/POS).",
            "",
        ],
        [
            "Aug 2026",
            "Ch. 4 SDLC",
            "Add a backlog to the development.",
            "Complied: PharmaLocate_Sprint_Documentation.docx (Steps 1–8 backlog) and PharmaLocate_Progress_Report.docx.",
            "",
        ],
        [
            "Aug 2026",
            "Ch. 4 Storyboard",
            "Add the desktop view. Show all features.",
            "Complied in the system at desktop width: Guest Home/Pharmacies/Medicines/Inquiries/Help; Auth; Admin dashboard, inquiries, inventory, POS, pharmacies, geofences, users, settings, backup. Paper should caption desktop storyboard figures.",
            "",
        ],
        [
            "Sep 2026",
            "References / figures",
            "Remove the border in the references. Add a figure name for each figure.",
            "Manuscript formatting: unboxed references; Figure x.x Title on every diagram and screenshot.",
            "",
        ],
        [
            "Sep 2026",
            "Geofencing",
            "Specify the parameter of the geofencing.",
            "Complied: Pharmacies tab “Geofencing parameters” card — circular (haversine), default center TPH 15.4870, 120.5960, default radius 5,000 m, admin range 500–10,000 m. Admin Geofences form repeats the same limits. TPH-internal pharmacies excluded.",
            "",
        ],
        [
            "Sep 2026",
            "Users / priority groups",
            "Make it general instead of focusing on senior citizens, pregnant women, PWDs, etc.",
            "Complied: priority_type / is_priority removed (July 2026). Home and Help copy say all users. Inquiries ordered by date, not by special group. Roles remain customer, staff, admin.",
            "",
        ],
        [
            "Sep 2026",
            "Devices",
            "Consider the concept of the devices.",
            "Complied: viewport meta; desktop-first layout; stacks at ≤800 px; Home pill “Desktop & mobile browsers.” Ordinary Chrome/Edge; no kiosk-only UI.",
            "",
        ],
        [
            "Oct 2026",
            "Respondents",
            "Break down respondents in a table. Interviews at the pharmacies mentioned. Insert evaluation-rating table.",
            "Field + manuscript: interview SpaRx and Magic 8 on-site; tabulate who / how many / role; insert Likert/evaluation table (may map to ISO/IEC 25010 in Documentation v1.3).",
            "",
        ],
        [
            "Oct 2026",
            "System close-out (adviser check)",
            "Verify that recommended features exist in the running prototype before pre-defense.",
            "Live at http://127.0.0.1:8000. Demo: maria / admin / sparx_staff (password). PHPUnit Sprint 7: 29/29; TPH-exclusion test passed. Remaining: binder screenshots; validated Help FAQs; manuscript tables/figures listed above.",
            "",
        ],
    ]
    data_table(
        doc,
        ["DATE", "SECTION / PART", "RECOMMENDATIONS", "STATUS / COMPLIANCE (team)", "SIGNATURE"],
        rows,
        [2.2, 3.2, 5.0, 5.4, 2.0],
    )

    p = doc.add_paragraph()
    r = p.add_run(
        "Overall technical-adviser finding as of 1 October 2026: Product implementation "
        "(inquiry, geofencing, inventory, POS) is complete for the agreed prototype scope. "
        "Manuscript items (SDG wording, short intro, IPOO figure, LLM cleanup, respondent tables, "
        "figure captions) remain for the paper. The technical adviser may sign the last column after verifying the live demo."
    )
    set_run(r, size=9)

    sign_block(
        doc,
        "Chua, Aurelio, Mercado, Lacanlale",
        "Proponents / researchers",
        "Kwinno L. Pineda",
        "Technical Adviser",
    )
    footer_line(doc, "TSU-CCS-SF-15     TCP Form 4 Rev. 2")
    path = OUT_DIR / "CCS TCP FORM4 rev 2(progress report-tech ad).docx"
    doc.save(path)
    return path


def make_form5():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.5)
        s.right_margin = Cm(1.5)
        s.page_width = Cm(21.59)
        s.page_height = Cm(27.94)

    header_block(doc, "TCP FORM 5", "Rev. 2", "DENNIS VIRTUDAZO", "SUBJECT TEACHER")
    meta_table(
        doc,
        [
            ("Project title", "PharmaLocate: A Web-Based Pharmacy Inquiry System with Geofencing for Real-Time Drug Availability, Sale & Inventory Checking"),
            ("Proponents", "Chua, Calvin O.; Aurelio, Lester C.; Mercado, Arji Chesner G.; Lacanlale, Nathalie Angel C."),
            ("Period covered", "July 2026 – 1 October 2026"),
            ("Purpose of this form", "Chronological deliverables submitted for subject-teacher monitoring, with the teacher’s recommendation or remark on each slice."),
        ],
    )
    p = doc.add_paragraph()
    r = p.add_run(
        "Instructions: DATE is when the slice was presented as done. DELIVERABLE is what the team produced. "
        "RECOMMENDATIONS are the subject teacher’s action (accepted, revise, or verify in class). "
        "SIGNATURE is left blank for the subject teacher."
    )
    set_run(r, size=9, color=MUTED)

    rows = [
        [
            "Jul 2026",
            "Step 1 — Frontend lock. Three-file UI (HTML/CSS/JS) frozen: Guest, Auth, and Admin screens; 10-medicine demo; tabs and sidebar as the screen contract. Close-out: 100%.",
            "Accepted as UI lock. Proceed to live API. Do not keep changing layout while connecting the database.",
            "",
        ],
        [
            "Jul 2026",
            "Step 2 — Backend map. Screen → functional requirement → API plan; P1–P5 order agreed (customer live data = P1; geofencing = P3; POS = P4; testing = P5). Close-out: 100%.",
            "Accepted as planning deliverable. Keep geofencing after stock is live.",
            "",
        ],
        [
            "Jul 2026",
            "Step 3 — Live API (P1). Laravel, MySQL, and Sanctum connected to the frozen UI: availability, pharmacies, inquiries. Guests see live stock without an account; customer maria can log in and send inquiries. Close-out: 100%.",
            "Accepted. Verify on http://127.0.0.1:8000, not the offline HTML demo file.",
            "",
        ],
        [
            "Jul 2026",
            "Step 4 — Admin core (P2). Live dashboard, inquiry replies, and stock edits. Staff limited to assigned pharmacy; /api/admin blocked for customers; admin/staff view opens after login. Priority user types (senior, PWD, pregnant, parent) removed; roles admin, staff, customer remain. Close-out: 100%.",
            "Accepted. Align later with panel instruction to treat users generally. Confirm staff cannot edit another store.",
            "",
        ],
        [
            "Jul 2026",
            "Sprint 5.0 — Baseline verification only (no new code): XAMPP, php artisan serve, login, live tabs before map work. Close-out: 100%.",
            "Accepted as process checkpoint.",
            "",
        ],
        [
            "Jul 2026",
            "Sprint 5.1 — Pharmacy CRUD API and admin form (create, edit, deactivate; GPS and hours). Close-out: 100%.",
            "Accepted. Profiles must stay community pharmacies around TPH, not inside TPH.",
            "",
        ],
        [
            "Jul 2026",
            "Sprint 5.2 — Geofence API, public zones, seeded Tarlac Provincial Hospital zone (5 km). Close-out: 100%.",
            "Accepted. Document radius and center for the paper.",
            "",
        ],
        [
            "Jul 2026",
            "Sprint 5.3 — Leaflet/OSM customer map replacing SVG; GPS optional with hospital default; CSS overflow fix. Close-out: 100%.",
            "Accepted. Demo with GPS denied is valid.",
            "",
        ],
        [
            "Jul 2026",
            "Sprint 5.4 — Server-side geofence filter (research core): only pharmacies in the containing zone; empty-zone notice. Close-out: 100%.",
            "Accepted as the geofencing contribution. Do not describe it as “pins only.”",
            "",
        ],
        [
            "Aug 2026",
            "Sprint 5.5 — Admin geofence UI (create/edit/assign) and admin map reusing existing APIs. Close-out: 100%.",
            "Accepted.",
            "",
        ],
        [
            "Aug 2026",
            "Sprint 5.6 — Optional Tile38 with haversine fallback so class demo works without Redis. Close-out: 100% with fallback.",
            "Accepted. /api/health may show tile38 unavailable.",
            "",
        ],
        [
            "Aug 2026",
            "Step 6 — POS and extras. POS decrements stock (oversell rejected; admin may choose pharmacy, staff cannot); user management (last admin protected); settings; JSON/CSV export. Close-out: 100%.",
            "Accepted. Title and 1.3 must name sale and inventory. Show a sale then the Medicines tab.",
            "",
        ],
        [
            "Aug 2026",
            "Sprints 7.0–7.4 — ISO/IEC 25010 testing. Plan in Documentation v1.3 (BB, WB, API IDs). PHPUnit 29/29 (August re-run). API-01–API-20 all pass. 62 of 80 mapped cases verified (~78%). Execution report issued.",
            "Accepted automation. Complete remaining black-box screenshots for the binder before final defense.",
            "",
        ],
        [
            "Sep 2026",
            "Sprint 8 — Polish and Help. Boot loader, tab motion, deferred GPS. Rule-based Help (no AI), 18 unofficial bilingual presets. Bug fixes: navbar name, fake inquiry samples, Help/Close toggle. Close-out: ~95%; validated FAQs pending.",
            "Accepted as polish. Replace unofficial FAQs when the validated list is ready. Help is not medical advice.",
            "",
        ],
        [
            "Oct 2026",
            "Panel (TCP Form 12) compliance in the running system: recommended title (availability, sale, inventory); TPH-internal pharmacies excluded; geofence parameters on the Pharmacies tab; inventory/POS labels; general-user copy; desktop and mobile layout. Tracker: Panel_Recommendations_Deliverables.docx.",
            "For verification in pre-defense. Run php artisan migrate. Subject teacher to confirm live demo against Form 12.",
            "",
        ],
    ]
    data_table(
        doc,
        ["DATE", "DELIVERABLE", "RECOMMENDATIONS (subject teacher)", "SIGNATURE"],
        rows,
        [2.2, 7.2, 5.6, 2.8],
    )

    p = doc.add_paragraph()
    r = p.add_run(
        "Summary for the subject teacher: Steps 1–6 and UI polish are complete for the prototype scope "
        "(2 pharmacies, 10 medicines, 1 geofence). Testing is largely automated; binder photos and the thesis "
        "narrative (SDG, IPOO, respondents) are the remaining academic deliverables. Recommended in-class check: "
        "guest map → maria inquiry → admin reply → POS sale. Demo accounts (password: password): maria (customer), "
        "admin (administrator), sparx_staff (SpaRx)."
    )
    set_run(r, size=9)

    sign_block(
        doc,
        "Chua, Aurelio, Mercado, Lacanlale",
        "Proponents / researchers",
        "Dennis Virtudazo",
        "Subject Teacher",
    )
    footer_line(doc, "TSU-CCS-SF-16     TCP Form 5 Rev. 2")
    path = OUT_DIR / "CCS TCP FORM5-rev2(progress report-subject teacher).docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    f4 = make_form4()
    f5 = make_form5()
    print(f4)
    print(f5)
