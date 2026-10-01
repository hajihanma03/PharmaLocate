# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "Panel_Recommendations_Status.docx"
GREEN = RGBColor(0x0F, 0x6E, 0x56)
DARK = RGBColor(0x1A, 0x26, 0x20)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def cell_text(cell, text, *, bold=False, color=DARK, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.bold = bold
    r.font.color.rgb = WHITE if white else color
    r.font.name = "Calibri"


def main():
    doc = Document()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PharmaLocate — TCP Form 12 recommendations")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = GREEN
    p = doc.add_paragraph()
    p.add_run("Source: CCS TCP FORM12 – recoss ng panels.pdf. Date applied in system: 1 October 2026.").font.size = Pt(11)

    rows = [
        ["Title — include POS; recommended wording", "System", "Browser title, navbar, and login subtitle now include availability, sale & inventory with geofencing."],
        ["Chapter 1 — stronger problem / SDG", "Manuscript", "Keep in the paper. Not a software change."],
        ["Introduction too long", "Manuscript", "Shorten Chapter 1 text in the document."],
        ["Add scope; no pharmacies inside TPH", "System", "Home scope banner; public API hides inside_tph pharmacies; admin checkbox."],
        ["Objective 2 format; inventory in 1.3", "System + paper", "Admin sidebar: Inventory management and POS — sales. Align 1.3 wording in the manuscript."],
        ["White-box testing as a table", "Paper + system", "WB cases already in Doc v1.3 and PHPUnit. Put the table in Chapter 1 as required."],
        ["Remove own study from LLM", "Manuscript", "Literature chapter only."],
        ["Revise conceptual framework IPOO", "Manuscript", "Diagram only."],
        ["Define features in terms", "Manuscript", "Use system feature names: inquiry, geofence, POS, inventory."],
        ["Architecture block diagram per icon", "Manuscript", "Draw icons; live architecture is Laravel + 3-file UI."],
        ["Sources of data ↔ objectives", "Manuscript", "Interview/pharmacy premises remains field work."],
        ["SDLC backlog", "System docs", "Sprint Documentation + Progress Report are the backlog log."],
        ["Storyboard desktop + all features", "System", "Desktop layout kept; home/pharmacies/medicines/inquiries + admin inventory/POS/geofences visible."],
        ["References border / figure names", "Manuscript", "Formatting in Word/LaTeX."],
        ["Specify geofencing parameters", "System", "Pharmacies tab and admin geofence form: circle, lat/lng, 500–10000 m, default 5 km at 15.4870, 120.5960."],
        ["General users, not seniors/PWD/etc.", "System", "Priority types already removed. Copy now says all users."],
        ["Consider devices", "System", "Desktop-first CSS; stacks on ≤800 px; scope pill lists desktop & mobile browsers."],
        ["Respondents table / interviews at pharmacies", "Field + paper", "Conduct interviews at SpaRx and Magic 8; insert rating table in the manuscript."],
    ]

    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Recommendation", "Where", "How it was satisfied"]):
        c = table.rows[0].cells[i]
        shade(c, "1D9E75")
        cell_text(c, h, bold=True, white=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = table.rows[i + 1].cells[j]
            if i % 2:
                shade(c, "E8F5EE")
            cell_text(c, val)

    n = doc.add_paragraph()
    n.add_run("Live database: run php artisan migrate in Programming4/backend so inside_tph exists on MySQL.").font.size = Pt(11)

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
